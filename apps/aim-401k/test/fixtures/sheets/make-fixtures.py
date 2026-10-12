#!/usr/bin/env python3
"""Regenerates the spreadsheet fixtures used by test/sheet.test.ts and scripts/e2e.d/40-sheet.mjs.

    python3 test/fixtures/sheets/make-fixtures.py

job-tracker.xlsx is written with openpyxl so it has the shape Excel produces
(shared strings, styles with date formats, booleans, a HYPERLINK() formula, a
cell hyperlink, a merged title row, a hidden sheet).
"""
import datetime as dt
import os
import re
import zipfile

import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))

# A personal tracker exported from Google Sheets (File > Download > CSV):
# a title row and a blank row above the header, quoted notes with commas,
# "" escapes and an embedded newline, a bare-domain link, a duplicate row.
GOOGLE = (
    "My job hunt 2026 (shared with Mom),,,,,\n"
    ",,,,,\n"
    "Company,Role,Link,Status,Date Applied,Notes\n"
    'Acme Robotics,Senior Platform Engineer,https://acme.example/careers/platform,Interviewing,10/02/2026,"Recruiter Wile E. said ""expect a call, next week"""\n'
    "Stark Industries,Frontend Engineer,https://jobs.lever.co/stark/fe-001?utm_source=linkedin,Applied,9/28/2026,\n"
    'Wayne Enterprises,Site Reliability Engineer,boards.greenhouse.io/wayne/jobs/555,Rejected,9/15/2026,"Two rounds,\nthen a form rejection"\n'
    "Cyberdyne Systems,ML Engineer,,Ghosted,8/20/2026,\n"
    "Soylent,Product Engineer,https://soylent.example/jobs/12,Wishlist,,Apply before Nov 1\n"
    "Stark Industries,Frontend Engineer,https://jobs.lever.co/stark/fe-001,Applied,9/28/2026,same posting twice\n"
    ",,,,,\n"
)

# A Notion database export: BOM, CRLF, "Month D, YYYY" dates, a Name title column.
NOTION = "\r\n".join(
    [
        "Name,Company,Position,Status,Applied,URL,Salary,Contact,Location,Created time",
        'Initrode - Staff SWE,Initrode,Staff Software Engineer,Applied,"October 3, 2026",https://initrode.example/jobs/staff,"$180,000 - $200,000",jane@initrode.example,"Austin, TX","October 1, 2026 9:15 AM"',
        'Massive Dynamic - EM,Massive Dynamic,Engineering Manager,Interviewing,"September 21, 2026",https://massivedynamic.example/careers/em,,Nina Sharpe (recruiter),Remote,"September 20, 2026 8:00 AM"',
        'Pied Piper - Backend,Pied Piper,Backend Engineer,Not started,,https://piedpiper.example/jobs/backend,,,"Palo Alto, CA","October 5, 2026 4:12 PM"',
        'Hooli - Platform,Hooli,Platform Engineer,Offer,"July 2, 2026",https://hooli.example/jobs/9,$210k,,Remote,"June 30, 2026 10:00 AM"',
        'Aviato,Aviato,Founding Engineer,Rejected,"August 12, 2026",,,,,"August 10, 2026 1:00 PM"',
    ]
) + "\r\n"


def write_text(name: str, text: str, bom: bool = False) -> None:
    with open(os.path.join(HERE, name), "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as f:
        f.write(text)


def excel_serial(d: dt.date) -> int:
    return (d - dt.date(1899, 12, 30)).days


def write_xlsx(name: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws["A1"] = "Job Search 2026"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    header = ["Company", "Job Title", "Status", "Date Applied", "Link", "Salary", "Remote?", "Notes"]
    ws.append([])  # row 2 blank
    ws.append(header)  # row 3
    for c in ws[3]:
        c.font = Font(bold=True)
    rows = [
        ["Globex Corporation", "Staff Engineer", "Phone Screen", dt.date(2026, 9, 14), "https://boards.greenhouse.io/globex/jobs/4242", 185000, True, "Referred by Hank"],
        ["Initech", "Backend Developer", "Rejected", dt.date(2026, 8, 30), "Posting", None, False, None],
        ["Umbrella Health", "Data Engineer", "Wishlist", None, '=HYPERLINK("https://umbrella.example/careers/77","apply here")', None, True, "Deadline Oct 30"],
        ["Hooli", "Platform Engineer", "Offer", dt.date(2026, 7, 2), "https://hooli.example/jobs/9", 210000, False, "Offer letter in Drive"],
        # Same job the core e2e flow submitted; the CSV import moves it to Interviewing first,
        # so this "Applied" row must not move it back.
        ["Acme Robotics", "Senior Platform Engineer", "Applied", dt.date(2026, 10, 1), "https://acme.example/careers/platform", None, False, None],
        # A date typed as a plain number (no date format): read by header + serial range.
        ["Black Mesa", "Research Engineer", "Assessment", excel_serial(dt.date(2026, 9, 1)), "https://blackmesa.example/jobs/3", None, False, "Take-home due Friday"],
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=4, max_row=3 + len(rows), min_col=4, max_col=4):
        for c in row:
            if isinstance(c.value, dt.date):
                c.number_format = "m/d/yyyy"
    # Cell hyperlink on Initech's "Posting" text (stored in the sheet rels).
    ws["E5"].hyperlink = "https://jobs.lever.co/initech/abc-123"
    ws["E5"].style = "Hyperlink"
    ws.column_dimensions["A"].width = 22

    archive = wb.create_sheet("Archive 2025")
    archive.append(["Company", "Role", "Stage", "Applied"])
    archive.append(["Vandelay Industries", "Importer/Exporter", "Ghosted", dt.date(2025, 3, 3)])
    archive.append(["Kramerica", "Intern", "Declined", dt.date(2025, 2, 1)])
    archive["D2"].number_format = "yyyy-mm-dd"
    archive["D3"].number_format = "yyyy-mm-dd"

    lists = wb.create_sheet("Lists")
    for s in ["Wishlist", "Applied", "Phone Screen", "Interview", "Offer", "Rejected"]:
        lists.append([s])
    lists.sheet_state = "hidden"

    path = os.path.join(HERE, name)
    wb.save(path)
    excelify(path)


def excelify(path: str) -> None:
    """openpyxl writes inline strings; Excel writes a shared-string table and caches
    formula results. Rewrite the package into Excel's shape (deflated parts,
    sharedStrings.xml with one rich-text item, cached HYPERLINK() result)."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    strings: list[str] = []
    index: dict[str, int] = {}

    def shared(m: "re.Match[str]") -> str:
        attrs, text = m.group(1), m.group(2)
        if text not in index:
            index[text] = len(strings)
            strings.append(text)
        return f'<c{attrs} t="s"><v>{index[text]}</v></c>'

    for n in sorted(parts):
        if re.match(r"xl/worksheets/sheet\d+\.xml$", n):
            xml = parts[n].decode()
            xml = re.sub(r'<c([^>]*?) t="inlineStr"><is><t>(.*?)</t></is></c>', shared, xml)
            xml = xml.replace(
                '<f>HYPERLINK("https://umbrella.example/careers/77","apply here")</f><v></v>',
                '<f>HYPERLINK("https://umbrella.example/careers/77","apply here")</f><v>apply here</v>',
            ).replace('<c r="E6"><f>', '<c r="E6" t="str"><f>')
            parts[n] = xml.encode()
    items = []
    for s in strings:
        if s == "Referred by Hank":  # rich text: two runs, the second bold
            items.append('<si><r><t xml:space="preserve">Referred by </t></r><r><rPr><b/><sz val="11"/></rPr><t>Hank</t></r></si>')
        else:
            items.append(f"<si><t>{s}</t></si>")
    parts["xl/sharedStrings.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(strings)}" uniqueCount="{len(strings)}">'
        + "".join(items)
        + "</sst>"
    ).encode()
    ct = parts["[Content_Types].xml"].decode()
    if "sharedStrings" not in ct:
        ct = ct.replace(
            "</Types>",
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        )
    parts["[Content_Types].xml"] = ct.encode()
    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    if "sharedStrings" not in rels:
        rels = rels.replace(
            "</Relationships>",
            '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="rIdSst"/></Relationships>',
        )
    parts["xl/_rels/workbook.xml.rels"] = rels.encode()
    order = ["[Content_Types].xml", "_rels/.rels"] + sorted(n for n in parts if n not in ("[Content_Types].xml", "_rels/.rels"))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n in order:
            info = zipfile.ZipInfo(n, date_time=(2026, 10, 1, 9, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, parts[n])


if __name__ == "__main__":
    write_text("google-tracker.csv", GOOGLE)
    write_text("notion-export.csv", NOTION, bom=True)
    write_xlsx("job-tracker.xlsx")
    print("wrote fixtures to", HERE)
